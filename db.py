"""Слой хранения данных системы учёта заказов.

Данные хранятся в рабочей книге Excel (``.xlsx``) — это файловое
хранилище, удовлетворяющее требованию ТЗ «хранение данных между
запусками (файлы или SQLite)». Класс :class:`Storage`:

- читает/пишет рабочую книгу ``store.xlsx`` с листами
  ``clients``, ``products``, ``orders``, ``order_items``, ``meta``;
- предоставляет CRUD-операции для клиентов, товаров и заказов;
- выполняет импорт и экспорт данных в форматах CSV и JSON;
- автоматически сохраняет изменения после каждой операции.

Дополнительно модуль содержит генератор демонстрационных данных
(:func:`sample_data`), позволяющий заполнить хранилище
реалистичными данными для визуализации и анализа.
"""

from __future__ import annotations

import csv
import json
import os
import random
import tempfile
from datetime import date, timedelta
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl
from openpyxl import Workbook, load_workbook

import models
from models import (
    Category,
    CorporateCustomer,
    Customer,
    Order,
    OrderItem,
    OrderStatus,
    Product,
    RegularCustomer,
    ValidationError,
    VIPCustomer,
)


class StorageError(Exception):
    """Ошибка работы с хранилищем данных.

    Возникает при повреждении файла, ошибках чтения/записи
    рабочей книги или нарушении целостности данных.
    """


#: Колонки листов рабочей книги. Порядок важен для сериализации.
SHEET_COLUMNS: Dict[str, List[str]] = {
    "clients": ["id", "type", "name", "email", "phone", "city", "registered", "orders_count", "inn", "priority"],
    "products": ["id", "name", "price", "category", "stock"],
    "orders": ["id", "customer_id", "customer_name", "customer_type", "order_date", "status", "total", "item_count"],
    "order_items": ["order_id", "product_id", "product_name", "price", "quantity"],
}


class Storage:
    """Хранилище данных на основе рабочей книги Excel.

    Parameters
    ----------
    path : str | Path
        Путь к файлу рабочей книги (``.xlsx``). Если файл
        отсутствует, создаётся новая пустая книга.
    """

    def __init__(self, path: str | Path = "data/store.xlsx") -> None:
        self.path: Path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.products: List[Product] = []
        self.customers: List[Customer] = []
        self.orders: List[Order] = []
        self._counters: Dict[str, int] = {"products": 1, "clients": 1, "orders": 1}
        self.load()

    # ------------------------------------------------------------------
    # Загрузка и сохранение рабочей книги
    # ------------------------------------------------------------------
    def load(self) -> None:
        """Загрузить данные из рабочей книги.

        Если файл отсутствует — создаётся пустая книга.
        При повреждении файла поднимается :class:`StorageError`.

        Raises
        ------
        StorageError
            Если рабочую книгу невозможно прочитать.
        """
        self.products.clear()
        self.customers.clear()
        self.orders.clear()
        self._counters = {"products": 1, "clients": 1, "orders": 1}

        if not self.path.exists():
            self._save()
            return

        try:
            wb = load_workbook(self.path, read_only=False, data_only=True)
        except Exception as exc:
            raise StorageError(f"Не удалось прочитать файл {self.path}: {exc}") from exc

        try:
            self._read_meta(wb)
            self._read_products(wb)
            self._read_customers(wb)
            self._read_orders(wb, wb["order_items"] if "order_items" in wb.sheetnames else None)
        finally:
            wb.close()
        self._recompute_counters()

    def _read_meta(self, wb: Workbook) -> None:
        """Прочитать счётчики идентификаторов с листа ``meta``."""
        if "meta" not in wb.sheetnames:
            return
        for row in wb["meta"].iter_rows(min_row=2, values_only=True):
            if row and row[0] in self._counters and row[1] is not None:
                self._counters[str(row[0])] = int(row[1])

    def _read_products(self, wb: Workbook) -> None:
        """Прочитать товары с листа ``products``."""
        if "products" not in wb.sheetnames:
            return
        for row in wb["products"].iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            product = Product(
                name=str(row[1]),
                price=float(row[2]),
                stock=int(row[4] or 0),
            )
            product.id = int(row[0])
            if row[3]:
                product.category = self._lookup_category(str(row[3]))
            self.products.append(product)

    def _read_customers(self, wb: Workbook) -> None:
        """Прочитать клиентов с листа ``clients``."""
        if "clients" not in wb.sheetnames:
            return
        for row in wb["clients"].iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            try:
                customer = Customer.from_dict(
                    {
                        "id": int(row[0]),
                        "type": str(row[1] or "Customer"),
                        "name": str(row[2]),
                        "email": str(row[3]),
                        "phone": str(row[4]),
                        "city": str(row[5] or ""),
                        "registered": str(row[6]) if row[6] else None,
                        "orders_count": int(row[7] or 0),
                        "inn": str(row[8] or ""),
                        "priority": int(row[9] or 1),
                    }
                )
            except (ValidationError, ValueError) as exc:
                raise StorageError(f"Повреждённая запись клиента (строка {row[0]}): {exc}") from exc
            self.customers.append(customer)

    def _read_orders(self, wb: Workbook, items_ws) -> None:
        """Прочитать заказы и связанные позиции."""
        if "orders" not in wb.sheetnames:
            return
        items_by_order: Dict[int, List[tuple]] = {}
        if items_ws is not None:
            for row in items_ws.iter_rows(min_row=2, values_only=True):
                if not row or row[0] is None:
                    continue
                items_by_order.setdefault(int(row[0]), []).append(row)

        customers_by_id = {c.id: c for c in self.customers}
        for row in wb["orders"].iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            order_id = int(row[0])
            customer = customers_by_id.get(int(row[1]))
            if customer is None:
                raise StorageError(f"Заказ #{order_id} ссылается на несуществующего клиента")
            try:
                order = Order(customer=customer, order_date=str(row[4]))
                order.id = order_id
                order.status = OrderStatus(str(row[5]))
            except (ValidationError, ValueError) as exc:
                raise StorageError(f"Повреждённый заказ #{order_id}: {exc}") from exc
            for item_row in items_by_order.get(order_id, []):
                product = Product(name=str(item_row[2]), price=float(item_row[3]))
                product.id = int(item_row[1]) if item_row[1] is not None else None
                order.add_item(product, quantity=int(item_row[4]))
            self.orders.append(order)

    def _lookup_category(self, name: str) -> Optional[Category]:
        """Поиск категории по имени (без дерева — плоский список)."""
        return Category(name)

    def _recompute_counters(self) -> None:
        """Пересчитать счётчики идентификаторов по максимуму."""
        for key, items in (("products", self.products), ("clients", self.customers), ("orders", self.orders)):
            ids = [obj.id for obj in items if obj.id is not None]
            self._counters[key] = max(ids) + 1 if ids else 1

    def _save(self) -> None:
        """Сохранить текущее состояние данных в рабочую книгу.

        Запись выполняется через временный файл с последующей
        заменой, чтобы не повредить книгу при сбое.

        Raises
        ------
        StorageError
            Если запись в файл не удалась.
        """
        wb = Workbook()
        ws_meta = wb.active
        ws_meta.title = "meta"
        ws_meta.append(["name", "value"])
        for key, value in self._counters.items():
            ws_meta.append([key, value])

        ws = wb.create_sheet("clients")
        ws.append(SHEET_COLUMNS["clients"])
        for customer in self.customers:
            data = customer.to_dict()
            ws.append(
                [
                    data.get("id"),
                    data.get("type"),
                    data.get("name"),
                    data.get("email"),
                    data.get("phone"),
                    data.get("city"),
                    data.get("registered"),
                    data.get("orders_count", 0),
                    data.get("inn", ""),
                    data.get("priority", ""),
                ]
            )

        ws = wb.create_sheet("products")
        ws.append(SHEET_COLUMNS["products"])
        for product in self.products:
            data = product.to_dict()
            ws.append(
                [
                    data.get("id"),
                    data.get("name"),
                    data.get("price"),
                    data.get("category", ""),
                    data.get("stock", 0),
                ]
            )

        ws = wb.create_sheet("orders")
        ws.append(SHEET_COLUMNS["orders"])
        ws_items = wb.create_sheet("order_items")
        ws_items.append(SHEET_COLUMNS["order_items"])
        for order in self.orders:
            data = order.to_dict()
            ws.append(
                [
                    data.get("id"),
                    data.get("customer_id"),
                    data.get("customer_name"),
                    data.get("customer_type"),
                    data.get("order_date"),
                    data.get("status"),
                    data.get("total"),
                    order.item_count(),
                ]
            )
            for item in order.items:
                idata = item.to_dict()
                ws_items.append(
                    [
                        data.get("id"),
                        idata.get("product_id"),
                        idata.get("product_name"),
                        idata.get("price"),
                        idata.get("quantity"),
                    ]
                )

        try:
            tmp_fd, tmp_path = tempfile.mkstemp(suffix=".xlsx", dir=self.path.parent)
            os.close(tmp_fd)
            wb.save(tmp_path)
            os.replace(tmp_path, self.path)
        except Exception as exc:
            raise StorageError(f"Не удалось сохранить файл {self.path}: {exc}") from exc
        finally:
            wb.close()

    # ------------------------------------------------------------------
    # CRUD: товары
    # ------------------------------------------------------------------
    def add_product(self, product: Product) -> Product:
        """Добавить товар в хранилище.

        Parameters
        ----------
        product : Product
            Добавляемый товар (без идентификатора).

        Returns
        -------
        Product
            Товар с назначенным идентификатором.
        """
        product.id = self._counters["products"]
        self._counters["products"] += 1
        self.products.append(product)
        self._save()
        return product

    def get_product(self, product_id: int) -> Optional[Product]:
        """Найти товар по идентификатору."""
        return next((p for p in self.products if p.id == product_id), None)

    def update_product(self, product: Product) -> None:
        """Обновить товар по идентификатору."""
        for idx, existing in enumerate(self.products):
            if existing.id == product.id:
                self.products[idx] = product
                self._save()
                return
        raise StorageError(f"Товар с id={product.id} не найден")

    def delete_product(self, product_id: int) -> None:
        """Удалить товар по идентификатору."""
        self.products = [p for p in self.products if p.id != product_id]
        self._save()

    # ------------------------------------------------------------------
    # CRUD: клиенты
    # ------------------------------------------------------------------
    def add_customer(self, customer: Customer) -> Customer:
        """Добавить клиента в хранилище."""
        customer.id = self._counters["clients"]
        self._counters["clients"] += 1
        self.customers.append(customer)
        self._save()
        return customer

    def get_customer(self, customer_id: int) -> Optional[Customer]:
        """Найти клиента по идентификатору."""
        return next((c for c in self.customers if c.id == customer_id), None)

    def update_customer(self, customer: Customer) -> None:
        """Обновить клиента по идентификатору."""
        for idx, existing in enumerate(self.customers):
            if existing.id == customer.id:
                self.customers[idx] = customer
                self._save()
                return
        raise StorageError(f"Клиент с id={customer.id} не найден")

    def delete_customer(self, customer_id: int) -> None:
        """Удалить клиента и связанные с ним заказы."""
        self.customers = [c for c in self.customers if c.id != customer_id]
        self.orders = [o for o in self.orders if o.customer.id != customer_id]
        self._save()

    # ------------------------------------------------------------------
    # CRUD: заказы
    # ------------------------------------------------------------------
    def add_order(self, order: Order) -> Order:
        """Добавить заказ в хранилище."""
        order.id = self._counters["orders"]
        self._counters["orders"] += 1
        self.orders.append(order)
        order.customer.orders_count += 1
        self._save()
        return order

    def get_order(self, order_id: int) -> Optional[Order]:
        """Найти заказ по идентификатору."""
        return next((o for o in self.orders if o.id == order_id), None)

    def update_order(self, order: Order) -> None:
        """Обновить заказ по идентификатору."""
        for idx, existing in enumerate(self.orders):
            if existing.id == order.id:
                self.orders[idx] = order
                self._save()
                return
        raise StorageError(f"Заказ с id={order.id} не найден")

    def delete_order(self, order_id: int) -> None:
        """Удалить заказ по идентификатору."""
        order = self.get_order(order_id)
        if order is None:
            raise StorageError(f"Заказ с id={order_id} не найден")
        self.orders = [o for o in self.orders if o.id != order_id]
        order.customer.orders_count = max(0, order.customer.orders_count - 1)
        self._save()

    # ------------------------------------------------------------------
    # Импорт / экспорт
    # ------------------------------------------------------------------
    def export_csv(self, entity: str, dest_dir: str | Path) -> Path:
        """Экспортировать сущность в CSV.

        Parameters
        ----------
        entity : str
            Одна из сущностей: ``clients``, ``products``, ``orders``.
        dest_dir : str | Path
            Каталог, куда будет сохранён файл.

        Returns
        -------
        Path
            Путь к созданному файлу.

        Raises
        ------
        StorageError
            При ошибке записи файла или неизвестной сущности.
        """
        return self._export_rows(entity, dest_dir, "csv")

    def export_json(self, entity: str, dest_dir: str | Path) -> Path:
        """Экспортировать сущность в JSON."""
        return self._export_rows(entity, dest_dir, "json")

    def _export_rows(self, entity: str, dest_dir: str | Path, fmt: str) -> Path:
        dest_dir = Path(dest_dir)
        dest_dir.mkdir(parents=True, exist_ok=True)
        data = self._entity_data(entity)
        columns = SHEET_COLUMNS.get(entity)
        if columns is None:
            raise StorageError(f"Неизвестная сущность: {entity}")
        name = f"{entity}.{fmt}"
        path = dest_dir / name
        try:
            if fmt == "csv":
                with open(path, "w", newline="", encoding="utf-8-sig") as fh:
                    writer = csv.DictWriter(fh, fieldnames=columns)
                    writer.writeheader()
                    writer.writerows(data)
            else:
                with open(path, "w", encoding="utf-8") as fh:
                    json.dump(data, fh, ensure_ascii=False, indent=2)
        except OSError as exc:
            raise StorageError(f"Не удалось экспортировать {entity} в {path}: {exc}") from exc
        return path

    def _entity_data(self, entity: str) -> List[dict]:
        """Получить плоские словари сущности для экспорта."""
        if entity == "clients":
            return [c.to_dict() for c in self.customers]
        if entity == "products":
            return [p.to_dict() for p in self.products]
        if entity == "orders":
            return [o.to_dict() for o in self.orders]
        raise StorageError(f"Неизвестная сущность: {entity}")

    def import_csv(self, entity: str, file_path: str | Path) -> int:
        """Импортировать данные сущности из CSV-файла.

        Импортированные записи добавляются к существующим
        (слияние с назначением новых идентификаторов).

        Parameters
        ----------
        entity : str
            Сущность: ``clients``, ``products`` или ``orders``.
        file_path : str | Path
            Путь к CSV-файлу.

        Returns
        -------
        int
            Количество импортированных записей.
        """
        file_path = Path(file_path)
        try:
            with open(file_path, newline="", encoding="utf-8-sig") as fh:
                rows = list(csv.DictReader(fh))
        except OSError as exc:
            raise StorageError(f"Не удалось открыть {file_path}: {exc}") from exc
        return self._import_rows(entity, rows)

    def import_json(self, entity: str, file_path: str | Path) -> int:
        """Импортировать данные сущности из JSON-файла."""
        file_path = Path(file_path)
        try:
            with open(file_path, encoding="utf-8") as fh:
                rows = json.load(fh)
        except (OSError, json.JSONDecodeError) as exc:
            raise StorageError(f"Не удалось прочитать JSON {file_path}: {exc}") from exc
        if not isinstance(rows, list):
            raise StorageError("JSON должен содержать список записей")
        return self._import_rows(entity, rows)

    def _import_rows(self, entity: str, rows: List[dict]) -> int:
        """Импортировать записи сущности (внутренняя логика)."""
        if entity == "clients":
            for row in self._clean_rows(rows, SHEET_COLUMNS["clients"]):
                try:
                    customer = Customer.from_dict(row)
                except (ValidationError, KeyError, TypeError, ValueError) as exc:
                    raise StorageError(f"Некорректная запись клиента: {exc}") from exc
                customer.id = None
                self.add_customer(customer)
            return len(rows)
        if entity == "products":
            for row in self._clean_rows(rows, SHEET_COLUMNS["products"]):
                try:
                    product = Product.from_dict(row)
                except (ValidationError, KeyError, TypeError, ValueError) as exc:
                    raise StorageError(f"Некорректная запись товара: {exc}") from exc
                product.id = None
                self.add_product(product)
            return len(rows)
        if entity == "orders":
            for row in rows:
                if not isinstance(row, dict):
                    raise StorageError("Каждая запись заказа должна быть словарём")
                try:
                    customer = self._resolve_customer(row)
                    order = Order.from_dict(row, customer)
                except (ValidationError, KeyError, TypeError, ValueError) as exc:
                    raise StorageError(f"Некорректная запись заказа: {exc}") from exc
                order.id = None
                self.add_order(order)
            return len(rows)
        raise StorageError(f"Неизвестная сущность: {entity}")

    @staticmethod
    def _clean_rows(rows: List[dict], columns: List[str]) -> List[dict]:
        """Отфильтровать словари, оставив известные колонки."""
        cleaned = []
        for row in rows:
            if not isinstance(row, dict):
                continue
            cleaned.append({key: row.get(key) for key in columns if key in row})
        return cleaned

    def _resolve_customer(self, order_row: dict) -> Customer:
        """Найти клиента заказа по ``customer_id`` либо по email."""
        customer_id = order_row.get("customer_id")
        if customer_id is not None:
            customer = self.get_customer(int(customer_id))
            if customer is not None:
                return customer
            raise StorageError(f"Клиент с id={customer_id} не найден")
        email = order_row.get("customer_email") or (order_row.get("customer", {}) or {}).get("email")
        for customer in self.customers:
            if customer.email == email:
                return customer
        raise StorageError("Не удалось определить клиента заказа (укажите customer_id)")

    # ------------------------------------------------------------------
    # Демонстрационные данные
    # ------------------------------------------------------------------
    def seed_demo_data(self, seed: int = 42) -> None:
        """Заполнить хранилище демонстрационными данными.

        Параметры
        ---------
        seed : int
            Зерно генератора случайных чисел для воспроизводимости.
        """
        customers, products, orders, _ = sample_data(seed)
        self.products, self.customers, self.orders = products, customers, orders
        self._recompute_counters()
        self._save()


def sample_data(seed: int = 42) -> Tuple[List[Customer], List[Product], List[Order], List[Category]]:
    """Сгенерировать демонстрационный набор данных.

    Parameters
    ----------
    seed : int, optional
        Зерно генератора случайных чисел (по умолчанию 42).

    Returns
    -------
    Tuple[List[Customer], List[Product], List[Order], List[Category]]
        Кортеж из клиентов, товаров, заказов и категорий.
    """
    rng = random.Random(seed)

    electronics = Category("Электроника")
    computers = Category("Компьютеры", electronics)
    smartphones = Category("Смартфоны", electronics)
    audio = Category("Аудио", electronics)
    home = Category("Дом и кухня")
    sport = Category("Спорт")

    product_specs = [
        ("Ноутбук ASUS VivoBook 15", 54990, computers),
        ("Ноутбук Lenovo IdeaPad 3", 42990, computers),
        ("Монитор Dell 27\"", 23990, computers),
        ("Клавиатура Logitech K380", 3990, computers),
        ("Мышь беспроводная Xiaomi", 1690, computers),
        ("iPhone 15 128 ГБ", 79990, smartphones),
        ("Samsung Galaxy A54", 31990, smartphones),
        ("Xiaomi Redmi Note 13", 21990, smartphones),
        ("Наушники Sony WH-1000XM5", 28990, audio),
        ("Колонка JBL Charge 5", 14990, audio),
        ("Микрофон Fifine K669", 3490, audio),
        ("Чайник Bosch TWK 3A011", 4590, home),
        ("Кофемашина DeLonghi", 45990, home),
        ("Пылесос Samsung SC4520", 7990, home),
        ("Микроволновка LG MS-2042", 9990, home),
        ("Утюг Tefal FV3940", 3490, home),
        ("Велосипед Stels Navigator", 23990, sport),
        ("Гантели 2×10 кг", 4990, sport),
        ("Массажный коврик", 1290, sport),
        ("Скакалка профессиональная", 890, sport),
    ]
    products = [Product(name, price, category, stock=rng.randint(5, 60)) for name, price, category in product_specs]
    for idx, product in enumerate(products, start=1):
        product.id = idx

    first_names = [
        "Иван", "Мария", "Алексей", "Елена", "Дмитрий", "Ольга", "Сергей", "Анна",
        "Николай", "Наталья", "Павел", "Юлия", "Андрей", "Екатерина", "Максим", "Татьяна",
        "Владимир", "Светлана", "Артём", "Ирина", "Роман", "Ксения", "Михаил", "Дарья",
        "Григорий", "Полина", "Виктор", "Вера", "Станислав", "Людмила",
    ]
    last_names = [
        "Иванов", "Петров", "Сидоров", "Смирнов", "Кузнецов", "Попов", "Васильев",
        "Павлов", "Соколов", "Михайлов", "Новиков", "Фёдоров", "Морозов", "Волков",
        "Алексеев", "Лебедев", "Семёнов", "Егоров", "Козлов", "Степанов",
    ]
    cities = ["Москва", "Санкт-Петербург", "Казань", "Новосибирск", "Екатеринбург", "Нижний Новгород", "Самара"]

    customers: List[Customer] = []
    for idx in range(1, 31):
        name = f"{rng.choice(first_names)} {rng.choice(last_names)}"
        email = f"user{idx:02d}@example.com"
        phone = f"+7 (9{rng.randint(0, 9)}{rng.randint(0, 9)}) {rng.randint(100, 999)}-{rng.randint(10, 99)}-{rng.randint(10, 99)}"
        city = rng.choice(cities)
        registered = (date(2024, 1, 1) + timedelta(days=rng.randint(0, 800))).isoformat()
        kind = rng.choices(
            ["regular", "vip", "corporate"], weights=[65, 25, 10], k=1
        )[0]
        if kind == "vip":
            customer: Customer = VIPCustomer(name, email, phone, city, registered, priority=rng.randint(1, 3))
        elif kind == "corporate":
            inn = "".join(str(rng.randint(0, 9)) for _ in range(10))
            customer = CorporateCustomer(name, email, phone, city, registered, inn=inn)
        else:
            customer = RegularCustomer(name, email, phone, city, registered)
        customer.id = idx
        customers.append(customer)

    orders: List[Order] = []
    end = date.today()
    start = end - timedelta(days=150)
    order_id = 1
    for customer in customers:
        n_orders = rng.randint(1, 8)
        for _ in range(n_orders):
            order_date = start + timedelta(days=rng.randint(0, (end - start).days))
            order = Order(customer, order_date.isoformat())
            n_items = rng.randint(1, 4)
            for _ in range(n_items):
                product = rng.choice(products)
                order.add_item(product, quantity=rng.randint(1, 3))
            order.status = rng.choice(list(OrderStatus))
            order.id = order_id
            order_id += 1
            customer.orders_count += 1
            orders.append(order)

    return customers, products, orders, [electronics, home, sport]